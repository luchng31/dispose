from __future__ import annotations

import csv
import io
import re
from typing import Any

from django.utils import timezone
from django.utils.crypto import get_random_string
from openpyxl import load_workbook

from apps.accounts.models import DeptLeaderMap, Role, User
from apps.assets.dispatch import dispatch_tickets_for_ip, remap_owner
from apps.assets.models import Asset, AssetOwnerMap

HEADER_ALIASES: dict[str, str] = {
    "ip": "ip", "主机ip": "ip", "ip地址": "ip", "资产ip": "ip", "内网ip": "ip",
    "hostname": "hostname", "主机名": "hostname", "资产名称": "hostname",
    "os": "os", "操作系统": "os",
    "biz_system": "biz_system", "业务系统": "biz_system", "业务": "biz_system",
    "dept": "biz_system", "部门": "biz_system",
    "owner": "owner", "负责人": "owner", "责任人": "owner",
    "管理人": "owner", "owner_username": "owner",
    "owner_dept": "owner_dept", "负责人部门": "owner_dept", "管理人-隶属组织": "owner_dept",
    "资源使用部门": "use_dept", "use_dept": "use_dept",
    "部门负责人": "dept_leader", "dept_leader": "dept_leader",
    "email": "email", "邮箱": "email",
    "wecom_userid": "wecom_userid", "企微账号": "wecom_userid", "工号": "wecom_userid",
}

# 负责人邮箱表（独立小表）：负责人/姓名(工号)/工号 + 邮箱
EMAIL_HEADER_ALIASES: dict[str, str] = {
    "owner": "owner", "负责人": "owner", "管理人": "owner", "姓名": "owner", "name": "owner",
    "emp_id": "emp_id", "工号": "emp_id", "wecom_userid": "emp_id", "企微账号": "emp_id",
    "email": "email", "邮箱": "email",
}

REQUIRED_CANON = ("ip",)
# owner name OR wecom 工号 either satisfies the owner requirement (工号-only
# sheets are auto-bound: username = 工号 = wecom_userid).
OWNER_CANON_OPTIONS = ("owner", "wecom_userid")

# 服务器资源汇总表.xlsx 列：内网IP*, 管理人*, 管理人-隶属组织, 资源使用部门, 部门负责人
TEMPLATE_HEADER = "内网IP,管理人,管理人-隶属组织,资源使用部门,部门负责人\n"
# 负责人邮箱表模板
EMAIL_TEMPLATE_HEADER = "负责人,工号,邮箱\n"

# 「解童钧(12047)」/「解童钧（12047）」-> ("解童钧", "12047")
_PERSON_RE = re.compile(r"^(.*?)\s*[（(]([^()（）]+)[）)]\s*$")


def normalize_header(name: str) -> str | None:
    return HEADER_ALIASES.get(str(name).strip().lower().replace(" ", ""))


def normalize_email_header(name: str) -> str | None:
    return EMAIL_HEADER_ALIASES.get(str(name).strip().lower().replace(" ", ""))


def parse_person(cell: str) -> tuple[str, str | None]:
    """Split a「姓名(工号)」cell into (name, emp_id); plain text -> (text, None)."""
    text = str(cell or "").strip()
    match = _PERSON_RE.match(text)
    if match is None:
        return text, None
    name = match.group(1).strip()
    emp_id = match.group(2).strip()
    return (name or text), (emp_id or None)


def _canon_row(raw: dict[str, str]) -> dict[str, str]:
    canon: dict[str, str] = {}
    for key, value in raw.items():
        hit = normalize_header(str(key))
        if hit is not None:
            canon[hit] = str(value or "").strip()
    return canon


def _canon_email_row(raw: dict[str, str]) -> dict[str, str]:
    canon: dict[str, str] = {}
    for key, value in raw.items():
        hit = normalize_email_header(str(key))
        if hit is not None:
            canon[hit] = str(value or "").strip()
    return canon


def _missing_required_columns(header_names: list[str]) -> list[dict[str, Any]]:
    canon = {c for c in (normalize_header(h) for h in header_names) if c is not None}
    missing: list[str] = []
    if "ip" not in canon:
        missing.append("ip")
    if not any(c in canon for c in OWNER_CANON_OPTIONS):
        missing.append("owner")
    if missing:
        return [{"row": "file", "message": f"缺少必需列：{','.join(missing)}"}]
    return []


def _missing_email_columns(header_names: list[str]) -> list[dict[str, Any]]:
    canon = {c for c in (normalize_email_header(h) for h in header_names) if c is not None}
    missing: list[str] = []
    if "email" not in canon:
        missing.append("邮箱")
    if "owner" not in canon and "emp_id" not in canon:
        missing.append("负责人(或工号)")
    if missing:
        return [{"row": "file", "message": f"缺少必需列：{','.join(missing)}"}]
    return []


def _resolve_user(name: str, emp_id: str | None) -> User | None:
    """Locate an existing user: by 工号 (wecom_userid) first, then username.

    A username hit whose 工号 differs from the incoming one is a different
    person (same display name) -> treated as not found so a distinct account
    gets created instead of silently merging two identities.
    """
    if emp_id:
        hit = User.objects.filter(wecom_userid=emp_id).first()
        if hit is not None:
            return hit
    if not name:
        return None
    candidate = User.objects.filter(username=name).first()
    if (
        candidate is not None
        and emp_id is not None
        and candidate.wecom_userid
        and candidate.wecom_userid != emp_id
    ):
        return None
    return candidate


def _create_user(
    name: str,
    emp_id: str | None,
    dept: str,
    role: str,
    created_users: list[dict[str, str]],
    email: str = "",
) -> User:
    # Same-name different-工号 people coexist: suffix the 工号 into the username.
    username = name
    if User.objects.filter(username=name).exists():
        username = f"{name}({emp_id})" if emp_id else f"{name}-{len(created_users)}"
    temp_password = get_random_string(12)
    user = User(
        username=username,
        role=role,
        is_active=True,
        dept=dept,
        email=email,
        wecom_userid=emp_id,
    )
    user.set_password(temp_password)
    user.save()
    created_users.append({"username": username, "temp_password": temp_password, "role": role})
    return user


def _provision_owner(
    name: str, emp_id: str | None, canon: dict[str, str], created_users: list[dict[str, str]]
) -> User:
    """Resolve/create the asset owner (管理人). Existing users keep their role/dept."""
    wecom = canon.get("wecom_userid", "").strip() or emp_id
    user = _resolve_user(name, wecom)
    if user is None:
        return _create_user(
            name, wecom, canon.get("owner_dept", ""), Role.OWNER, created_users,
            email=canon.get("email", ""),
        )
    touched = False
    if canon.get("email") and not user.email:
        user.email = canon["email"]
        touched = True
    if wecom and not getattr(user, "wecom_userid", ""):
        user.wecom_userid = wecom
        touched = True
    if touched:
        user.save(update_fields=["email", "wecom_userid"])
    return user


def _provision_leader(
    name: str,
    emp_id: str | None,
    dept_path: str,
    created_users: list[dict[str, str]],
    summary: dict[str, Any],
) -> User:
    """Resolve/create the department head (部门负责人) as Role.LEADER.

    Dept binding itself is handled by the caller-collected pairs (rebuilt or
    get_or_create'd after the row loop).
    """
    user = _resolve_user(name, emp_id)
    if user is None:
        user = _create_user(name, emp_id, dept_path, Role.LEADER, created_users)
        summary["created_leaders"] += 1
    elif user.role == Role.OWNER:
        # A manager who also heads a department gets promoted; never demote others.
        user.role = Role.LEADER
        user.save(update_fields=["role"])
        summary["upgraded_leaders"].append(user.username)
    updates: list[str] = []
    if not user.dept:
        user.dept = dept_path
        updates.append("dept")
    if emp_id and not user.wecom_userid:
        user.wecom_userid = emp_id
        updates.append("wecom_userid")
    if updates:
        user.save(update_fields=updates)
    return user


def _apply_one(
    raw: dict[str, str],
    idx: int,
    summary: dict[str, Any],
    imported_ips: set[str],
    leader_pairs: set[tuple[int, str]],
) -> None:
    canon = _canon_row(raw)
    ip = canon.get("ip", "")
    owner_name, owner_emp = parse_person(canon.get("owner", ""))
    wecom_fallback = canon.get("wecom_userid", "").strip()
    if not owner_name and wecom_fallback:
        # 工号-only rows keep the legacy binding: username = 工号 = wecom_userid.
        owner_name, owner_emp = wecom_fallback, None
    if not ip or not owner_name:
        summary["errors"].append({"row": idx, "message": "缺少 内网IP 或 管理人"})
        return
    # Only non-empty cells overwrite the asset, so sheets that lack hostname/os
    # columns never blank out existing values.
    defaults = {
        key: canon[key]
        for key in ("hostname", "os", "biz_system")
        if canon.get(key, "")
    }
    if canon.get("use_dept", ""):
        defaults["biz_system"] = canon["use_dept"]
    asset, created = Asset.objects.update_or_create(ip=ip, defaults=defaults)
    summary["created_assets" if created else "updated_assets"] += 1
    user = _provision_owner(owner_name, owner_emp, canon, summary["created_users"])
    before = asset.owner_history.filter(valid_to=None).first()
    remap_owner(asset, user)
    after = asset.owner_history.filter(valid_to=None).first()
    if (before.id if before else None) != (after.id if after else None):
        summary["remapped"] += 1
    summary["dispatched"] += dispatch_tickets_for_ip(ip)
    imported_ips.add(ip)

    leader_cell = canon.get("dept_leader", "")
    if leader_cell:
        leader_name, leader_emp = parse_person(leader_cell)
        if leader_name and canon.get("use_dept", ""):
            leader = _provision_leader(
                leader_name, leader_emp, canon["use_dept"], summary["created_users"], summary
            )
            leader_pairs.add((leader.pk, canon["use_dept"]))
        elif leader_name:
            summary["errors"].append(
                {"row": idx, "message": "部门负责人缺少资源使用部门，未建立部门映射"}
            )


def apply_asset_rows(rows: list[dict[str, str]], sync: bool = False) -> dict[str, Any]:
    """Upsert assets + owner map + provision users + dispatch.

    sync=True means the sheet is the authoritative full table (最新导入为准):
    IPs absent from the sheet get their current owner map closed (asset stays,
    becomes ownerless, open tickets re-dispatch to the orphan pool) and
    DeptLeaderMap is rebuilt from the sheet. sync=False is additive-only.
    """
    summary: dict[str, Any] = {
        "created_assets": 0, "updated_assets": 0, "created_users": [],
        "created_leaders": 0, "upgraded_leaders": [], "leader_dept_maps": 0,
        "leader_maps_rebuilt": False, "orphaned": 0,
        "remapped": 0, "dispatched": 0, "errors": [],
        "imported_at": timezone.now().isoformat(),
    }
    imported_ips: set[str] = set()
    leader_pairs: set[tuple[int, str]] = set()
    for idx, raw in enumerate(rows, start=2):
        _apply_one(raw, idx, summary, imported_ips, leader_pairs)
    if sync:
        now = timezone.now()
        stale: list[str] = list(
            AssetOwnerMap.objects.filter(valid_to=None)
            .exclude(ip_id__in=imported_ips)
            .values_list("ip_id", flat=True)
        )
        summary["orphaned"] = AssetOwnerMap.objects.filter(
            valid_to=None, ip_id__in=stale
        ).update(valid_to=now)
        for ip in stale:
            summary["dispatched"] += dispatch_tickets_for_ip(ip)
        DeptLeaderMap.objects.all().delete()
        DeptLeaderMap.objects.bulk_create(
            [DeptLeaderMap(user_id=uid, dept_path=path) for uid, path in sorted(leader_pairs)]
        )
        summary["leader_maps_rebuilt"] = True
        summary["leader_dept_maps"] = len(leader_pairs)
    else:
        for user_id, dept_path in leader_pairs:
            _, created = DeptLeaderMap.objects.get_or_create(user_id=user_id, dept_path=dept_path)
            if created:
                summary["leader_dept_maps"] += 1
    return summary


def is_summary_format(rows: list[dict[str, str]]) -> bool:
    """True when the sheet carries the 服务器资源汇总表 authoritative columns."""
    canon: set[str] = set()
    for row in rows[:1]:
        canon.update(c for c in (normalize_header(k) for k in row) if c is not None)
    return "use_dept" in canon or "dept_leader" in canon


def count_sync_orphans(rows: list[dict[str, str]]) -> int:
    """How many currently-owned IPs would become ownerless under full sync."""
    ips = {
        canon.get("ip", "")
        for row in rows
        if (canon := _canon_row(row)).get("ip", "")
    }
    return AssetOwnerMap.objects.filter(valid_to=None).exclude(ip_id__in=ips).count()


def apply_owner_email_rows(
    rows: list[dict[str, str]], commit: bool = True
) -> dict[str, Any]:
    """Update existing users' emails from the owner-email sheet (no account creation).

    Rows whose 负责人/工号 match no user land in ``missing``; matched rows
    land in ``changes`` [{user_id, username, old, new}] so the view can audit.
    commit=False previews counts without writing (dry-run).
    """
    summary: dict[str, Any] = {
        "total_rows": 0, "updated": 0, "unchanged": 0,
        "changes": [], "missing": [], "errors": [],
    }
    for idx, raw in enumerate(rows, start=2):
        summary["total_rows"] += 1
        canon = _canon_email_row(raw)
        email = canon.get("email", "").strip()
        name, emp = parse_person(canon.get("owner", ""))
        emp = canon.get("emp_id", "").strip() or emp
        if not email or (not name and not emp):
            summary["errors"].append({"row": idx, "message": "缺少 邮箱 或 负责人/工号"})
            continue
        user = _resolve_user(name, emp)
        if user is None:
            summary["missing"].append({"row": idx, "name": name or emp or ""})
            continue
        if not commit:
            summary["updated"] += 1
            continue
        old = user.email
        if old == email:
            summary["unchanged"] += 1
            continue
        user.email = email
        user.save(update_fields=["email"])
        summary["updated"] += 1
        summary["changes"].append(
            {"user_id": user.pk, "username": user.username, "old": old, "new": email}
        )
    return summary


def parse_asset_csv(data: bytes) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    rows, errors = _parse_sheet(data)
    if errors:
        return [], errors
    header_names = list(rows[0].keys()) if rows else []
    errors = _missing_required_columns(header_names)
    if errors:
        return [], errors
    return rows, []


def _parse_sheet(data: bytes) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    """Parse an uploaded sheet (xlsx sniffed by magic bytes, else CSV)."""
    if data[:4] == b"PK\x03\x04":
        return _rows_from_xlsx(data)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("gbk")
        except UnicodeDecodeError:
            return [], [{"row": "file", "message": "文件编码无法识别（需 UTF-8/GBK 的 CSV 或 xlsx）"}]
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        return [], [{"row": "file", "message": "文件为空或无表头"}]
    rows = [
        {str(k): str(v or "").strip() for k, v in row.items()}
        for row in reader
        if any((v or "").strip() for v in row.values())
    ]
    return rows, []


def _rows_from_xlsx(data: bytes) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return [], [{"row": "file", "message": "无法解析 xlsx 文件（需合法的 Excel 工作簿）"}]
    try:
        worksheet = workbook.worksheets[0]
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            return [], [{"row": "file", "message": "表格为空或无表头"}]
        names = [str(h).strip() for h in header if h is not None and str(h).strip()]
        rows: list[dict[str, str]] = []
        for raw in iterator:
            row = {
                name: "" if value is None else str(value).strip()
                for name, value in (
                    (str(h).strip(), v) for h, v in zip(header, raw, strict=False)
                    if h is not None and str(h).strip()
                )
            }
            if any(row.get(name, "") for name in names):
                rows.append(row)
    finally:
        workbook.close()
    return rows, []


def parse_asset_file(filename: str, data: bytes) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    """Parse an asset+owner sheet: 服务器资源汇总表 xlsx or the legacy CSV dialect."""
    rows, errors = _parse_sheet(data)
    if errors:
        return [], errors
    header_names = list(rows[0].keys()) if rows else []
    errors = _missing_required_columns(header_names)
    if errors:
        return [], errors
    return rows, []


def parse_owner_email_file(filename: str, data: bytes) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    """Parse the owner-email sheet (负责人/工号 + 邮箱); .xlsx or CSV."""
    rows, errors = _parse_sheet(data)
    if errors:
        return [], errors
    header_names = list(rows[0].keys()) if rows else []
    errors = _missing_email_columns(header_names)
    if errors:
        return [], errors
    return rows, []
